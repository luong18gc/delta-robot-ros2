#!/usr/bin/env python3
"""
Khung định dạng dùng chung cho khóa luận bản 2 — mỗi chương một tệp, cùng một bộ quy tắc trình bày.

Tách ra khỏi `build_thesis_v2.py` khi viết thêm Chương 3–4: chép lại ba trăm dòng định dạng cho
mỗi tệp chương thì sửa một quy tắc phải sửa nhiều chỗ, và chắc chắn có chỗ bị bỏ sót, nên hai
chương sẽ dần trình bày khác nhau mà không ai cố ý.

Dùng:

    import thesis_style as S
    S.configure(OUT_PATH, FIGURES, TABLES)
    S.H1('CHUONG 3. ...')
    S.P('...')
    S.save()

Định dạng theo *Quy định về trình bày ĐATN* của Trường ĐH Công nghệ, ĐHQGHN: A4, Times New Roman
13 pt, dãn dòng 1,3; lề trên 2,5 / dưới 3 / trái 3 / phải 2 cm; tiêu đề bảng đặt TRÊN, tiêu đề
hình đặt DƯỚI.
"""

import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import omml as m                                                   # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, '..', 'figures')
OUT = None
FONT, BODY_PT = 'Times New Roman', 13
TEXT_W = Cm(16)          # bề rộng vùng chữ: A4 21 − lề trái 3 − lề phải 2

doc = Document()

FIGURES, TABLES = [], []
FIG_USED, TAB_USED = set(), set()
_FIG, _TAB = {}, {}


def configure(out_path, figures, tables):
    """Chọn tệp đầu ra và đăng ký danh mục hình/bảng của chương đang dựng."""
    global OUT, FIGURES, TABLES, _FIG, _TAB
    OUT, FIGURES, TABLES = out_path, figures, tables
    _FIG = {k: (cap, fn, w) for k, cap, fn, w in figures}
    _TAB = dict(tables)
    setup()


def save():
    """Ghi tệp, sau khi kiểm mọi hình/bảng đã khai báo đều được dùng đúng một lần."""
    missing_fig = {k for k, *_ in FIGURES} - FIG_USED
    missing_tab = {k for k, _ in TABLES} - TAB_USED
    assert not missing_fig, f'Hình khai báo nhưng chưa dùng: {sorted(missing_fig)}'
    assert not missing_tab, f'Bảng khai báo nhưng chưa dùng: {sorted(missing_tab)}'
    doc.save(OUT)
    print(f'Đã ghi {os.path.normpath(OUT)}')
    print(f'  {len(FIG_USED)} hình, {len(TAB_USED)} bảng')


# ------------------------------------------------------------------ định dạng
def _insert_ordered(parent, element, after_tags):
    for child in parent:
        if child.tag.split('}')[1] in after_tags:
            child.addprevious(element)
            return
    parent.append(element)


_AFTER_PGNUM = {'cols', 'formProt', 'vAlign', 'noEndnote', 'titlePg', 'textDirection',
                'bidi', 'rtlGutter', 'docGrid', 'printerSettings', 'footnotePr', 'endnotePr'}
_AFTER_SHD_TCPR = {'tcMar', 'textDirection', 'tcFitText', 'vAlign', 'hideMark'}


def page_setup(sec):
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(3)
    sec.left_margin, sec.right_margin = Cm(3), Cm(2)


def setup():
    page_setup(doc.sections[0])
    st = doc.styles['Normal']
    st.font.name = FONT
    st.font.size = Pt(BODY_PT)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    pf = st.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.3
    pf.first_line_indent = Cm(1)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for level, size in ((1, 15), (2, 14), (3, 13)):
        h = doc.styles[f'Heading {level}']
        h.font.name = FONT
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = RGBColor(0, 0, 0)
        h.font.italic = False
        h.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
        hp = h.paragraph_format
        hp.space_before = Pt(12 if level == 1 else 10)
        hp.space_after = Pt(6)
        hp.first_line_indent = Cm(0)
        hp.line_spacing = 1.3
        hp.keep_with_next = True
    for name in ('List Bullet', 'List Number'):
        s = doc.styles[name]
        s.font.name = FONT
        s.font.size = Pt(BODY_PT)
        s.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
        s.paragraph_format.space_after = Pt(3)
        s.paragraph_format.line_spacing = 1.3
    settings = doc.settings.element
    upd = OxmlElement('w:updateFields')
    upd.set(qn('w:val'), 'true')
    _insert_ordered(settings, upd, {
        'hdrShapeDefaults', 'footnotePr', 'endnotePr', 'compat', 'docVars', 'rsids',
        'mathPr', 'attachedSchema', 'themeFontLang', 'clrSchemeMapping'})
    for zoom in settings.iter(qn('w:zoom')):
        if zoom.get(qn('w:percent')) is None:
            zoom.set(qn('w:percent'), '100')
    sec = doc.sections[0]
    sec.footer.is_linked_to_previous = False
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run()
    for kind, text in (('begin', None), (None, 'PAGE'), ('separate', None), ('end', None)):
        if kind:
            el = OxmlElement('w:fldChar')
            el.set(qn('w:fldCharType'), kind)
        else:
            el = OxmlElement('w:instrText')
            el.set(qn('xml:space'), 'preserve')
            el.text = text
        run._r.append(el)


INLINE = re.compile(r'(\*\*.+?\*\*|\*.+?\*|`.+?`)')


def add_runs(p, text, size=None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2])
            r.bold = True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            r = p.add_run(part[1:-1])
            r.italic = True
        elif part.startswith('`') and part.endswith('`'):
            r = p.add_run(part[1:-1])
            r.font.name = 'Consolas'
        else:
            r = p.add_run(part)
        r.font.size = Pt(size or BODY_PT)
        r.font.name = r.font.name or FONT
    return p


def BREAK():
    doc.add_paragraph().add_run().add_break(6)


def H1(title, page_break=True):
    if page_break:
        doc.add_paragraph().runs or None
        doc.add_page_break()
    doc.add_heading(title, 1)


def H2(t):
    doc.add_heading(t, 2)


def H3(t):
    doc.add_heading(t, 3)


def P(text, indent=True, size=None):
    p = doc.add_paragraph()
    add_runs(p, text, size=size)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if not indent:
        p.paragraph_format.first_line_indent = Cm(0)
    return p


def B(items):
    for it in items:
        p = doc.add_paragraph(style='List Bullet')
        add_runs(p, it)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(0)


def N(items):
    for it in items:
        p = doc.add_paragraph(style='List Number')
        add_runs(p, it)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(0)


def EQ(body, number=None):
    """
    Công thức Word thật: căn GIỮA trang, số hiệu căn PHẢI lề.

    Dùng hai điểm dừng tab — một CENTER ở giữa vùng chữ, một RIGHT ở mép phải — thay vì gõ mấy dấu
    tab rồi canh bằng mắt. Nhờ vậy công thức luôn đúng giữa và số hiệu luôn đúng lề, không lệch khi
    nội dung dài ngắn khác nhau.
    """
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    pf.tab_stops.add_tab_stop(int(TEXT_W / 2), WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(TEXT_W, WD_TAB_ALIGNMENT.RIGHT)
    p.add_run('\t')
    p._p.append(parse_xml(f'<m:oMath {nsdecls("m")}>{body}</m:oMath>'))
    if number:
        r = p.add_run(f'\t({number})')
        r.font.size = Pt(BODY_PT)
        r.font.name = FONT
    return p


def IMG(key):
    """Hình: ảnh căn giữa, tiêu đề đặt DƯỚI theo quy định trình bày."""
    caption, filename, width = _FIG[key]
    assert key not in FIG_USED, f'Hình {key} dùng hai lần'
    FIG_USED.add(key)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(os.path.join(FIG, filename), width=Cm(width))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.first_line_indent = Cm(0)
    c.paragraph_format.space_after = Pt(12)
    add_runs(c, f'**Hình {key}.** {caption}', size=12)
    return c


def shade(cell, color):
    el = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color)
    _insert_ordered(el, shd, _AFTER_SHD_TCPR)


def T(key, header, rows, widths=None, size=12):
    """Bảng: tiêu đề đặt TRÊN theo quy định trình bày."""
    caption = _TAB[key]
    assert key not in TAB_USED, f'Bảng {key} dùng hai lần'
    TAB_USED.add(key)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    add_runs(p, f'**Bảng {key}.** {caption}', size=12)
    table = doc.add_table(rows=1, cols=len(header))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ''
        cp = cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.first_line_indent = Cm(0)
        cp.paragraph_format.space_after = Pt(2)
        add_runs(cp, f'**{h}**', size=size)
        shade(cell, 'E9EDF2')
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ''
            cp = cells[i].paragraphs[0]
            cp.paragraph_format.first_line_indent = Cm(0)
            cp.paragraph_format.space_after = Pt(2)
            cp.alignment = (WD_ALIGN_PARAGRAPH.LEFT if i == 0
                            else WD_ALIGN_PARAGRAPH.CENTER)
            add_runs(cp, str(val), size=size)
    if widths:
        for r in table.rows:
            for i, w in enumerate(widths):
                r.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


