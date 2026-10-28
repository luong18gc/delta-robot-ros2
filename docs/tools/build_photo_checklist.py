#!/usr/bin/env python3
"""
Dựng danh sách ảnh cần chụp TRƯỚC KHI tháo băng tải và camera.

    python3 docs/tools/build_photo_checklist.py

Chia theo tiêu chí DUY NHẤT có ý nghĩa lúc này: **chụp lại được hay không sau khi tháo**. Ảnh của
phần cứng và ảnh của camera thật thì không; ảnh của mô phỏng thì lúc nào cũng chụp được.
"""

import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'Danh_sach_anh_can_chup.docx')
FONT = 'Times New Roman'
doc = Document()


def setup():
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
    sec.left_margin, sec.right_margin = Cm(2.0), Cm(1.5)
    style = doc.styles['Normal']
    style.font.name = FONT
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    style.paragraph_format.space_after = Pt(4)


def H(text, size=15, color=None, before=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return p


def P(text, size=12, italic=False, color=None):
    p = doc.add_paragraph()
    for i, chunk in enumerate(text.split('**')):
        run = p.add_run(chunk)
        run.bold = i % 2 == 1
        run.italic = italic
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
    return p


def shade(cell, color):
    el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear')
    el.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(el)


def table(rows, widths, header_fill='1F3864'):
    t = doc.add_table(rows=1, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    head = ['', 'Ảnh cần chụp', 'Chụp thế nào', 'Dùng ở đâu']
    for cell, text, w in zip(t.rows[0].cells, head, widths):
        cell.width = Cm(w)
        shade(cell, header_fill)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor.from_string('FFFFFF')
    for row in rows:
        cells = t.add_row().cells
        for cell, text, w in zip(cells, ('☐',) + row, widths):
            cell.width = Cm(w)
            para = cell.paragraphs[0]
            if text == '☐':
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for i, chunk in enumerate(str(text).split('**')):
                run = para.add_run(chunk)
                run.bold = i % 2 == 1
                run.font.size = Pt(16 if text == '☐' else 10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


setup()

H('DANH SÁCH ẢNH CẦN CHỤP TRƯỚC KHI THÁO BĂNG TẢI VÀ CAMERA', 16, '1F3864', before=0)
P('Chia theo tiêu chí duy nhất có ý nghĩa lúc này: **sau khi tháo thì còn chụp lại được không**. '
  'Phần A và B là không; phần C thì lúc nào cũng chụp được nên để sau cũng kịp.')
P('Mẹo chung: chụp ngang (landscape), đủ sáng, lấy nét vào vật chính. Ảnh dùng cho khóa luận in '
  'đen trắng, nên tránh dựa vào màu để phân biệt — nếu cần thì dán nhãn giấy ghi chữ.', italic=True)

H('⚠ QUAN TRỌNG NHẤT: QUAY VIDEO TRƯỚC', 14, 'C00000')
P('Ảnh tĩnh không thay được video khi bảo vệ. Quay **ba đoạn**, mỗi đoạn 30–60 giây, để ngang:')
P('1. **Một vòng băng tải trọn vẹn** — đặt lon, băng chạy, robot bám theo, băng dừng, robot gắp '
  'và thả vào khay. Quay sao cho khung hình thấy **cả bàn thật và màn hình mô phỏng**.')
P('2. **Bản sao số** — dùng tay đẩy một lon trên bàn, quay màn hình Gazebo cho thấy lon ảo đi theo.')
P('3. **Hệ thống từ chối** — đưa tay vào khung hình, quay màn hình cho thấy hệ thống không bị '
  'đánh lừa (lon ảo đứng yên, nhật ký ghi lý do).')
P('Nếu chỉ kịp một đoạn thì chọn đoạn 1.', italic=True)

H('A. PHẦN CỨNG — tháo rồi là mất hẳn', 14, 'C00000')

table([
    ('**Toàn cảnh bố trí**: bàn, camera trên giá, băng tải, ba lon, sáu marker',
     'Đứng lùi 1,5–2 m, chụp chếch từ phía sau camera để thấy cả hướng nhìn',
     'Chương 3 — bố trí thực nghiệm'),
    ('**Camera C270 trên giá kẹp**, cận cảnh',
     'Thấy rõ cách gá vào mép bàn và hướng chúc xuống',
     'Chương 3 — bố trí camera'),
    ('**Đo khoảng cách camera**: đặt thước dây từ chân giá tới mép bàn, và thước đứng đo chiều cao',
     'Chụp sao cho **đọc được số trên thước**. Số thiết kế: lùi 557 mm, cao 631 mm',
     'Chương 3 — đối chứng cho tư thế camera đo bằng hiệu chuẩn'),
    ('**Băng tải toàn bộ**, nhìn từ trên xuống',
     'Thấy hai con lăn, mặt băng, chiều dài; đặt thước cạnh băng để có tỉ lệ',
     'Chương 3 — thông số băng tải'),
    ('**Cụm động cơ và khớp nối**, cận cảnh', 'Thấy động cơ, trục, cách truyền động',
     'Chương 3 — cơ cấu băng tải'),
    ('**Mạch điện**: Arduino + mạch cầu H + nguồn',
     'Chụp từ trên, thấy rõ các dây nối; có thể dán nhãn giấy ghi tên chân',
     'Chương 3 — sơ đồ điện (kèm hình vẽ)'),
    ('**Sáu marker trên bàn**, nhìn gần như từ trên xuống',
     'Thấy đủ ID 0–5 và vị trí tương đối; có thước càng tốt',
     'Chương 2 — hiệu chuẩn ngoại tham số'),
    ('**Mặt bàn TRỐNG** (không có lon, không có băng tải nếu tháo được)',
     'Cùng góc với camera, để minh họa ảnh nền dùng cho phép trừ nền',
     'Chương 2 — tách vật khỏi nền'),
    ('**Ba lon thật đứng cạnh nhau**', 'Chụp gần, thấy rõ nhãn và màu của từng lon',
     'Chương 1 hoặc 3 — đối tượng thí nghiệm'),
    ('**Lon trên băng tải ở đầu nạp** và **ở điểm gắp**',
     'Hai ảnh riêng, cùng một góc, để thấy quãng đường lon đi',
     'Chương 3 — vòng làm việc'),
    ('**Bố trí chiếu sáng**: đèn đặt ở đâu',
     'Chụp toàn cảnh có đèn trong khung', 'Chương 2 — điều kiện ánh sáng'),
    ('**Thước đặt cạnh một lon**', 'Đo đường kính và chiều cao lon thật',
     'Chương 3 — hệ số tỉ lệ k = 3'),
], widths=[1.0, 5.4, 6.2, 4.4])

H('B. ẢNH MÀN HÌNH CẦN CAMERA THẬT — tháo rồi là không chụp lại được', 14, 'C00000')

P('Chạy: **ros2 launch delta_controller conveyor_cycle.launch.py**, rồi mở ảnh gỡ lỗi bằng '
  '**ros2 run rqt_image_view rqt_image_view /vision/debug_image**. Chụp màn hình bằng phím '
  '**PrtSc** (hoặc Shift+PrtSc để chọn vùng).')

table([
    ('**Ba lon nhận dạng đủ**, cả ba đều tin cậy',
     'Đặt ba lon cách nhau ≥ 80 mm trên bàn, chụp ảnh gỡ lỗi',
     'Chương 2 và 4 — kết quả nhận dạng'),
    ('**Lon trên băng tải** đang được nhận dạng',
     'Đặt lon lên băng, chụp ảnh gỡ lỗi', 'Chương 3 — giả thuyết mặt đỡ'),
    ('**Bàn tay che một lon** → hệ thống từ chối',
     'Đưa tay che một phần lon, chụp lúc nhãn đổi sang cảnh báo',
     'Chương 2 — cờ tin cậy'),
    ('**Hai lon chạm nhau** → bị từ chối',
     'Đặt hai lon sát nhau, chụp ảnh gỡ lỗi', 'Chương 4 — giới hạn hệ thống'),
    ('**Sáu marker được nhận dạng**, ảnh có khoanh marker',
     'Chụp khi hệ thống báo 6/6 marker', 'Chương 2 — hiệu chuẩn'),
    ('**Ảnh nền** (mặt bàn trống qua camera)',
     'Mở tệp **calibration/table_reference.png** rồi chụp màn hình, hoặc chép thẳng tệp đó',
     'Chương 2 — phép trừ nền'),
    ('**So sánh ngưỡng màu và trừ nền** trên cùng một cảnh',
     'Nếu còn kịp: chạy **scripts/try_object_detector.py**; nếu không, dùng hình đã có',
     'Chương 2 — hai phương pháp tách vật'),
    ('**Bản sao số hai màn hình**: ảnh camera thật và cửa sổ Gazebo cạnh nhau',
     'Chạy **scripts/view_twin.py**, hoặc xếp hai cửa sổ cạnh nhau rồi chụp cả màn hình',
     'Chương 3 — bản sao số (ảnh quan trọng nhất của chương)'),
    ('**Robot bám lon trên băng**',
     'Chạy vòng băng tải, chụp lúc đầu hút đang ở trên nắp lon đang di chuyển',
     'Chương 3 — bám theo vật'),
    ('**Nhật ký vòng băng tải** trong cửa sổ dòng lệnh',
     'Chụp lúc thấy đủ các dòng: bám → chạy → dừng → gắp → chờ nhấc lon',
     'Chương 3 — máy trạng thái'),
    ('**Giao diện điều khiển** đang chạy, bảng vật thể có số liệu',
     'Chạy **ros2 run delta_controller gui**, chụp toàn cửa sổ',
     'Chương 3 — phần mềm điều khiển'),
], widths=[1.0, 5.4, 6.2, 4.4])

H('C. ẢNH CHỈ CẦN MÔ PHỎNG — để sau cũng được', 14, '2E7D32')

P('Những ảnh này chụp lại lúc nào cũng được, kể cả sau khi đã tháo hết phần cứng. **Không cần vội.**',
  italic=True)

table([
    ('Cảnh mô phỏng nhìn từ camera ảo, ba khay có nhãn',
     '**ros2 launch delta_controller pick_place.launch.py** rồi xem '
     '**/side_camera/image**', 'Chương 3 — cảnh mô phỏng'),
    ('Cửa sổ Gazebo nhìn tổng thể robot và cảnh', 'Chụp màn hình Gazebo',
     'Chương 3 — môi trường mô phỏng'),
    ('Robot đang gắp lon trong mô phỏng', 'Chạy **Phân loại tất cả**, chụp lúc đang hút',
     'Chương 3 — gắp thả'),
    ('Ba lon đã được phân loại vào ba khay', 'Chụp sau khi chạy xong',
     'Chương 4 — kết quả'),
    ('Ảnh gỡ lỗi của camera MÔ PHỎNG', 'Xem **/vision/debug_image**',
     'Chương 2 — nhận dạng theo màu'),
], widths=[1.0, 5.4, 6.2, 4.4])

H('D. Việc cần làm ngay sau khi chụp xong', 14, '1F3864')

P('1. Chép toàn bộ ảnh vào **docs/figures/** trong thư mục đồ án, đặt tên có nghĩa '
  '(ví dụ **anh_bo_tri_tong_the.jpg**, **anh_bang_tai_tren_xuong.jpg**).')
P('2. Chạy **git add -A && git commit && git push** để ảnh được sao lưu lên GitHub — phần cứng '
  'tháo rồi thì ảnh là thứ duy nhất còn lại.')
P('3. Chép luôn các tệp hiệu chuẩn hiện tại (**calibration/**) nếu chưa commit: chúng ghi lại tư '
  'thế camera thật, và không dựng lại được sau khi tháo.')

P('')
P('⚠ Nếu thiếu thời gian, thứ tự ưu tiên: **video đoạn 1 → mục A → mục B → mục C**.',
  color='C00000')

doc.save(OUT)
print(f'Đã ghi {os.path.normpath(OUT)}')
sys.exit(0)
